"""Build the review workbook from supplied semantics and the proposed page map."""
from pathlib import Path
from collections import Counter
import json
import re
import sys
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(sys.argv[1])
PAGES = json.loads((Path(__file__).with_name('seo-map-2026-10.json')).read_text(encoding='utf-8'))
APPLIED = all(p.get('status', '').startswith('Внедрено локально') for p in PAGES)
REPORTS = ROOT / '_reports'
PERIOD = '21.04.2026–20.05.2026'
SOURCE_NAME = SOURCE.name
raw = BeautifulSoup(SOURCE.read_text(encoding='utf-8'), 'html.parser')
queries = {}
for table in raw.find_all('table'):
    headers = [c.get_text(' ', strip=True) for c in table.find('tr').find_all(['td', 'th'])]
    if 'Запрос' not in headers:
        continue
    qi = headers.index('Запрос')
    for row in table.find_all('tr')[1:]:
        cells = [c.get_text(' ', strip=True) for c in row.find_all(['td', 'th'])]
        if len(cells) <= qi + 1:
            continue
        query = re.sub(r'\s+(конкурент|наш бренд)$', '', cells[qi])
        query = query.replace(' (общий запрос)', '')
        if not query or query.startswith('Сумма'):
            continue
        freq = cells[qi + 1]
        # Never convert the author's cluster estimates into measured frequencies.
        if query not in queries:
            queries[query] = {'freq': int(freq.replace(' ', '')) if re.fullmatch(r'[\d ]+', freq) else freq,
                              'source': SOURCE_NAME, 'period': PERIOD}

for page in PAGES:
    for query in page['queries'].split('; '):
        queries.setdefault(query, {'freq': None, 'source': 'Предложение к расширению семантики', 'period': None})

def classify(query):
    q = query.lower()
    if any(s in q for s in ['диорама', 'спецкомбинат', 'комбинат ритуальных', 'агентство память']):
        return '', 'Чужие бренды', 'Исключить', 'Навигационные запросы к другим компаниям.'
    if 'животных' in q:
        return '', 'Кремация животных', 'Исключить', 'Эта услуга не подтверждена и не относится к текущей странице кремации.'
    if 'прижизненный' in q:
        return '', 'Прижизненный договор', 'Уточнить услугу', 'Проверить наличие услуги и старый URL; новую страницу пока не создавать.'
    if 'сво' in q:
        return '', 'Похороны участников СВО', 'Уточнить условия', 'Подтвердить условия и порядок помощи; затем решить, нужен ли раздел в организации похорон.'
    if any(s in q for s in ['памятник', 'оградк', 'гроб', 'венок', 'цветы на похороны']):
        return '', 'Каталог и работы с памятниками', 'Отложено', 'Товарные запросы — будущие категории каталога. Установку памятника отдельно подтвердить; не смешивать продажу и монтаж.'
    if 'реквизит' in q or 'инн' in q:
        index = 10
    elif 'отзыв' in q:
        index = 9
    elif any(s in q for s in ['адрес', 'телефон', 'казарменный']):
        index = 8
    elif any(s in q for s in ['о компании', 'офис', 'зал прощания']):
        index = 7
    elif q in PAGES[6]['queries'].split('; '):
        index = 6
    elif 'кремаци' in q:
        index = 2
    elif 'благоустройств' in q or 'уход за могил' in q:
        index = 3
    elif any(s in q for s in ['груз', 'транспортиров', 'перевозк']):
        index = 4
    elif any(s in q for s in ['документ', 'пособие', 'юридическ']):
        index = 5
    elif 'похорон' in q:
        index = 1
    else:
        index = 0
    page = PAGES[index]
    note = ''
    if 'цены' in q or 'стоит' in q or 'стоимость' in q or 'цена' in q:
        note = 'Ценовой раздел этой страницы; фактические суммы добавить после получения прайса.'
    if q == 'ритуальные услуги хабаровск адреса':
        note = 'Вторичный запрос: выдача может ожидать список компаний. Основной фокус контактов — бренд + адрес/телефон.'
    if q == 'похороны':
        note = 'Слишком общий запрос: проверить коммерческий смысл и географию, исключить фильмы, песни и нерелевантные значения.'
    if q == 'пособие на похороны':
        note = 'Информационный подраздел услуги. Проверить актуальные условия по официальным источникам перед публикацией.'
    if 'уход за могил' in q:
        note = 'Пока подраздел благоустройства; самостоятельную услугу и страницу подтвердить отдельно.'
    source = queries[query]['source']
    status = 'Распределён из отчёта' if source == SOURCE_NAME else 'Новый вариант — проверить'
    if isinstance(queries[query]['freq'], str):
        status = 'Оценка автора отчёта — проверить'
    return page['url'], page['cluster'], status, note

wb = Workbook()
wb.remove(wb.active)
def sheet(name, headers, rows, widths, row_height=84):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'C2'
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='233D31')
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 34
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Calibri', size=11, color='292C29')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[row[0].row].height = row_height
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    tab = Table(displayName='Map'+str(len(wb.worksheets)), ref=ws.dimensions)
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=True)
    ws.add_table(tab)
    ws.sheet_view.zoomScale = 80
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:1'
    return ws

page_rows = [[p['priority'],p['page'],p['url'],p['cluster'],p['queries'],p['h1'],p['title'],len(p['title']),p['description'],len(p['description']),p['content'],p['note'],p.get('status','Черновик, в сайт не внесён')] for p in PAGES]
sheet('Карта страниц', ['Очередь','Страница','URL','Группа запросов','Основные и дополнительные запросы','H1','Title','Знаков Title','Description','Знаков Description','Что раскрыть на странице','Примечание','Статус H1 и метатегов'],page_rows,[9,25,44,35,60,48,55,12,75,15,70,70,30],150)

query_rows = []
for query, meta in queries.items():
    url, cluster, status, note = classify(query)
    query_rows.append([query,url,cluster,status,meta['freq'],meta['period'],meta['source'],note,None,None])
query_rows.sort(key=lambda row: (not bool(row[1]),row[1],row[0]))
sheet('Запросы', ['Запрос','Целевая страница','Группа','Статус','Частота из старого отчёта','Период старого замера','Источник','Примечание','Новая частота — заполнить','Регион, операторы и дата нового замера'],query_rows,[52,44,38,33,20,25,46,78,22,42],66)

deferred = [
 ['Каталог','Памятники, гробы, венки, ограды','Категории → товары → характеристики → обращение','Фотографии, ассортимент, материалы, размеры, цены и условия заказа. URL категорий пока не назначены.'],
 ['Наши работы','Реальные выполненные проекты','Страница проекта и связь с соответствующей услугой','Фотографии результата, состав работ, материалы и разрешение на публикацию. Выставочные образцы не подменяют проекты.'],
 ['Вопрос-ответ','Общие информационные вопросы','Тематические группы вопросов со ссылками на услуги','Ответы заказчика. Вопросы конкретной услуги пока раскрываются в её существующем FAQ.'],
 ['Статьи','Информационные запросы','Отдельные материалы с конкретной задачей читателя','Согласованный план и достоверные материалы. Сейчас раздел не наполняем.'],
 ['Прижизненный договор','прижизненный договор хабаровск','Возможная отдельная услуга','Подтверждение, что услуга оказывается, условия, документы и старый URL для переноса.'],
 ['Уход за захоронениями','уход за могилами хабаровск','Сейчас подраздел благоустройства','Для отдельной страницы нужны перечень работ, периодичность, условия заказа и расчёта.'],
 ['Кремация животных и чужие бренды','Нерелевантные текущему предложению запросы','Исключены','Не переносим в тексты и метатеги существующих страниц.'],
]
sheet('Отложено и исключено',['Раздел','Запросы или тема','Решение','Что нужно'],deferred,[30,50,62,90],80)

notes = [
 ['Назначение','Рабочая SEO-карта на 14.10.2026. ' + ('Согласованные H1, Title и Description внедрены на локальном сайте. Публикация на сервер отдельно.' if APPLIED else 'H1, Title и Description — предложения для согласования; в WordPress они не внесены.') + ' Таблица не предназначена для размещения на сайте.'],
 ['Состав',f'{len(PAGES)} существующих страниц и {len(query_rows)} уникальных формулировок из доступных таблиц отчёта и предложений к расширению. Это не полное семантическое ядро на 500 запросов из ТЗ.'],
 ['Очередь','1 — главная и пять услуг. 2 — список услуг, о компании, контакты, отзывы и реквизиты. Отложенные разделы ждут информации заказчика.'],
 ['Частоты','Источник: '+SOURCE_NAME+'. Выгрузка 25.05.2026 за '+PERIOD+'. Старые значения не описывают спрос октября. Приблизительные оценки автора сохранены как текст. Пустая частота означает отсутствие замера, а не нулевой спрос.'],
 ['География','В исходнике указано «регион 76 — Хабаровск (Хабаровский край)». Перед новым замером уточнить выбранный регион, операторы запроса и период. Суммировать пересекающиеся частоты как число людей нельзя.'],
 ['Правило распределения','Общие запросы об агентстве закреплены за главной; конкретная услуга и её стоимость — за страницей услуги. Одна формулировка имеет одну основную посадочную. Это редакционная карта, распределение уточняется по реальной выдаче и намерению посетителя.'],
 ['Метатеги','Не повторять все ключевые слова подряд. Title должен коротко и точно описывать страницу. Google может сформировать другой заголовок; жёсткого лимита символов нет. Диапазоны из ТЗ — редакционный ориентир, длины показаны отдельно.'],
 ['H1 и слоган','H1 главной объясняет вид услуги и город. ' + ('Слоган «Когда ночь, когда день» сохранён отдельной видимой строкой под H1.' if APPLIED else 'Слоган «Когда ночь, когда день» можно сохранить отдельной видимой строкой.')],
 ['Цены','Пока цены не получены, на сайте остаются …. Не включать в Description обещание опубликованного прайса или минимальную сумму. После получения данных — заполнить разделы цен соответствующих услуг.'],
 ['Следующий шаг',('H1 и метатеги внедрены и проверены. Обновить Wordstat и проверить выдачу, получить подтверждённые цены. Далее — формы, аналитика и подготовка переноса/индексации.' if APPLIED else 'Согласовать распределение и формулировки; обновить Wordstat и проверить выдачу; внедрить H1/Title/Description; проверить отображение и отсутствие дублей. Далее — формы, аналитика и подготовка переноса/индексации.')],
 ['Wordstat','https://yandex.ru/support2/wordstat/ru/interface/new'],
 ['Title Google','https://developers.google.com/search/docs/appearance/title-link'],
]
sheet('Как работать',['Тема','Пояснение'],notes,[28,125],88)

destination = REPORTS / ('seo-map-2026-10-14-applied.xlsx' if APPLIED else 'seo-map-2026-10-14.xlsx')
wb.save(destination)
check = load_workbook(destination, read_only=True)
assert check['Карта страниц'].max_row == len(PAGES)+1
assert check['Запросы'].max_row == len(query_rows)+1
assert len({p['url'] for p in PAGES}) == len(PAGES)
assert len({p['title'] for p in PAGES}) == len(PAGES)
check.close()

md = ['# Распределение поисковых запросов по страницам','',
      ('Согласованные H1, Title и Description внедрены на локальном сайте 14 октября 2026. Публикация на сервер не выполнялась.' if APPLIED else 'Черновик от 14 октября 2026. Метатеги и H1 пока не внесены в сайт.') + ' Стек остаётся прежним: WordPress и текущая тема.','',
      f'[Открыть полную таблицу Excel]({destination.name}). В ней четыре листа: карта страниц, отдельные запросы, отложенные темы и порядок работы.','',
      '## Основное распределение','',
      '| Страница | Группа запросов | Примеры |','|---|---|---|']
for p in PAGES:
    examples = '; '.join(p['queries'].split('; ')[:2])
    md.append(f"| [{p['page']}](http://ra-zil.test{p['url']}) | {p['cluster']} | {examples} |")
md += ['', '## Какие столбцы нужны', '',
       'На листе «Карта страниц»: очередь → страница → URL → группа запросов → основные и дополнительные запросы → H1 → Title → Description → что раскрыть → примечание → статус H1 и метатегов. Длина метатегов вынесена в отдельные столбцы. Статус внедрения не означает завершения всех задач по содержанию и замерам спроса.', '',
       'На листе «Запросы» — одна строка на запрос: целевая страница, группа, статус, старая частота и её период, источник, примечание, поля для нового замера. Запросы для проверки отделены от формулировок из исходного отчёта.', '',
       '## ' + ('Принятые решения' if APPLIED else 'Что согласовать перед внедрением'), '',
       '1. Общие запросы «ритуальные услуги / агентство Хабаровск» ведут на главную. Список услуг помогает выбрать направление. Каждая конкретная услуга и её стоимость раскрываются на своей странице.',
       ('2. На главной установлен H1 «Ритуальные услуги в Хабаровске», слоган «Когда ночь, когда день» сохранён отдельной строкой. Заголовки услуг приведены к согласованной карте.' if APPLIED else '2. Для главной предложен H1 «Ритуальные услуги в Хабаровске». Слоган «Когда ночь, когда день» можно сохранить отдельной видимой строкой. Для услуг предложены понятные заголовки с городом, без перечисления всех вариантов ключевых слов.'),
       '3. Уход за захоронениями пока остаётся частью благоустройства. Самостоятельную страницу ухода и прижизненного договора создаём после подтверждения услуг и содержания.',
       '4. Каталог, наши работы, общий вопрос-ответ и статьи остаются отложенными. Чужие бренды и кремация животных исключены из продвижения текущих услуг.', '',
       '## Как понимать частоты и метатеги', '',
       f'Количество уникальных формулировок — {len(query_rows)}. Из них из доступных таблиц исходника — {sum(m["source"] == SOURCE_NAME for m in queries.values())}, предложено для проверки — {sum(m["source"] != SOURCE_NAME for m in queries.values())}. Это рабочая карта текущих страниц, а не полное ядро на 500 запросов из ТЗ. Частоты перенесены из выгрузки 25 мая за {PERIOD}; они не являются свежим замером. Оценочные значения автора отмечены отдельно. Пустая ячейка означает, что частота не измерена.', '',
       'Регион в исходнике описан неоднозначно: Хабаровск и Хабаровский край. Для нового замера нужно фиксировать выбранную географию, период и операторы. Частоты пересекающихся запросов нельзя складывать в число уникальных посетителей. [Справка Wordstat](https://yandex.ru/support2/wordstat/ru/interface/new).', '',
       'Title предложены уникальными и описательными. Длина показана для редакционной проверки, но не является гарантией отображения: Google может сформировать другой заголовок и ограничивает его доступной шириной. [Рекомендации Google](https://developers.google.com/search/docs/appearance/title-link).', '',
       ('H1, Title и Description проверены в HTML всех 11 страниц и совпадают с картой. Следующие задачи — обновить замеры спроса, проверить намерение запросов по выдаче, настроить аналитику и подготовить перенос с проверкой индексации.' if APPLIED else 'Сначала согласуем карту и метатеги, обновим замеры и проверим намерение запросов по выдаче. Затем можно внедрять изменения в главную и пять услуг, после них — во вспомогательные страницы.') + ' Для ценовых запросов нужны реальные суммы от заказчика; текущие заглушки не заменяем выдуманными ценами.', '']
(REPORTS/'seo-map-2026-10-14.md').write_text('\n'.join(md),encoding='utf-8')
print(json.dumps({'pages':len(PAGES),'queries':len(query_rows),'statuses':dict(Counter(r[3] for r in query_rows)),'title_lengths':{p['page']:len(p['title']) for p in PAGES},'description_lengths':{p['page']:len(p['description']) for p in PAGES}},ensure_ascii=False))
